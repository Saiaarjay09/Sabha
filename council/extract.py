"""Turn an uploaded file into plain text.

Extraction happens in memory and the bytes are dropped as soon as the text is
out — nothing here writes to disk.
"""

from __future__ import annotations

import io
import re


class ExtractionError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ").replace("•", "- ").replace("●", "- ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def from_pdf(data: bytes) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ExtractionError("PDF support needs pypdf installed on the server.") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as e:
        raise ExtractionError(f"Could not read that PDF ({type(e).__name__}).") from e

    # Structural signals a real parser reacts to, which are lost the moment the
    # document is flattened to a string. Counting images matters because a CV
    # that puts its contact details or skills in a graphic is, to an ATS,
    # a CV that does not have them.
    images = 0
    for page in reader.pages:
        try:
            images += len(page.images)
        except Exception:
            pass
    signals = {"pages": len(reader.pages), "images": images}

    text = _clean("\n\n".join(pages))
    if len(text) < 120:
        # Almost certainly a scan or an image-only export. Worth saying so
        # precisely, because it is also exactly why an ATS would reject it.
        raise ExtractionError(
            "That PDF has almost no extractable text — it looks like a scan or an "
            "image-only export. An ATS would read it the same way, as blank. "
            "Export a text-based PDF, or paste the text instead."
        )
    return text, signals


def from_docx(data: bytes) -> tuple[str, dict]:
    try:
        import docx
    except ImportError as e:
        raise ExtractionError("DOCX support needs python-docx installed on the server.") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not read that Word file ({type(e).__name__}).") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:                      # CVs in tables are common and ATS-hostile
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    # These were previously read and silently discarded, which meant the audit
    # could never report the single most common reason a Word CV parses badly.
    try:
        xml = d.element.xml
    except Exception:
        xml = ""
    signals = {
        "tables": len(d.tables),
        # Text boxes are worse than tables: python-docx cannot read them and
        # neither can many parsers, so their content is invisible on both sides.
        # Count opening tags only — "txbxContent" matches the closing tag too.
        "textboxes": xml.count("<w:txbxContent"),
        # inline_shapes are already <w:drawing> elements, so counting both
        # double-counts every inline image. Take the larger of the two views.
        "images": max(len(d.inline_shapes), xml.count("<w:drawing")),
    }
    return _clean("\n".join(parts)), signals


def from_upload(filename: str, data: bytes) -> tuple[str, str, dict]:
    """Returns (text, source, signals) where source is pdf | docx | text.

    `signals` carries the structural facts that only exist while the file is
    still a file — tables, text boxes, embedded images, page count. The audit
    needs them because they are what actually breaks a parser, and they are
    unrecoverable once the document has been flattened to a string.
    """
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        text, sig = from_pdf(data)
        return text, "pdf", sig
    if name.endswith(".docx") or data[:2] == b"PK":
        text, sig = from_docx(data)
        return text, "docx", sig
    if name.endswith(".doc"):
        raise ExtractionError("Legacy .doc isn't supported — save as PDF or .docx, or paste the text.")
    try:
        return _clean(data.decode("utf-8")), "text", {}
    except UnicodeDecodeError:
        try:
            return _clean(data.decode("latin-1")), "text", {}
        except Exception as e:
            raise ExtractionError("Couldn't read that file as text.") from e
